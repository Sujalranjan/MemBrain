"""Builders for small, valid test documents."""

import io

import docx


def make_pdf(*pages: str) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]  # pages object filled in below
    page_ids = []
    font_id = 3
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for text in pages:
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("latin-1")
        stream = b"BT /F1 12 Tf 72 720 Td (" + escaped + b") Tj ET"
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        content_id = len(objects)
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
            b"/Resources << /Font << /F1 %d 0 R >> >> >>" % (content_id, font_id)
        )
        page_ids.append(len(objects))
    kids = b" ".join(b"%d 0 R" % i for i in page_ids)
    objects[1] = b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % len(page_ids)

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
        document.add_paragraph("After the table.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()

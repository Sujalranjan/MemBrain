"""Plain-text extraction from uploaded documents (PDF, DOCX, text formats)."""

import io
import re
import zipfile
from pathlib import PurePath

import docx
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from pypdf import PdfReader
from pypdf.errors import PdfReadError

FORMATS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "text",
    ".md": "text",
    ".markdown": "text",
    ".csv": "text",
    ".log": "text",
}


class UnsupportedDocument(ValueError):
    """The file type is not one we can extract text from."""


class ExtractionError(ValueError):
    """The file claims a supported type but its text could not be extracted."""


def detect_format(filename: str, data: bytes) -> str:
    """Pick the extractor from the extension, then check the file's signature actually matches it."""
    fmt = FORMATS.get(PurePath(filename).suffix.lower())
    if fmt is None:
        supported = ", ".join(sorted(ext.lstrip(".").upper() for ext in FORMATS))
        raise UnsupportedDocument(f"Unsupported file type; supported: {supported}")
    if fmt == "pdf" and not data.startswith(b"%PDF-"):
        raise ExtractionError("File has a .pdf extension but is not a PDF")
    if fmt == "docx" and not data.startswith(b"PK"):
        raise ExtractionError("File has a .docx extension but is not a Word document (legacy .doc is not supported)")
    return fmt


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        # Many PDFs are "encrypted" with an empty user password purely to set permissions.
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("PDF is password-protected")
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except (PdfReadError, ValueError, KeyError) as e:
        if isinstance(e, ExtractionError):
            raise
        raise ExtractionError(f"Could not read PDF: {e}") from e


def _docx_text(data: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as e:
        raise ExtractionError(f"Could not read Word document: {e}") from e
    blocks = []
    # Paragraphs and tables in document order, so table content stays next to the text describing it.
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            for row in item.rows:
                cells = list(dict.fromkeys(c.text.strip() for c in row.cells))  # merged cells repeat
                blocks.append(" | ".join(c for c in cells if c))
        else:
            blocks.append(item.text)
    return "\n\n".join(blocks)


def _plain_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def normalize(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_text(filename: str, data: bytes) -> str:
    fmt = detect_format(filename, data)
    text = normalize({"pdf": _pdf_text, "docx": _docx_text, "text": _plain_text}[fmt](data))
    if not text:
        hint = " (scanned PDFs need OCR, which is not supported)" if fmt == "pdf" else ""
        raise ExtractionError(f"No extractable text found{hint}")
    return text
